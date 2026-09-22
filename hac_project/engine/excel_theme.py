"""
ENGINE — Excel Theme (ถอดจาก 6N5.xlsx)
=======================================
โมดูลนี้เก็บเฉพาะ "หน้าตา" (ธีมสี Office ของ 6N5.xlsx, ฟอนต์ Browallia New, พาเลตสีต่อ role,
เส้นขอบ, ความกว้าง/สูง, number format) ที่ถอดมาจากไฟล์ต้นแบบ 6N5.xlsx จริง (ยืนยันโดยผู้ใช้เป็น
ต้นแบบ format 2026-09) — ไม่มีลำดับแถว/ข้อความ/สูตรคำนวณอะไรอยู่ในนี้เลย ของพวกนั้นทั้งหมด
ยังคงเป็นของเราเองใน engine/excel_export.py (ห้ามยืมจาก 6N5 แม้จะคล้ายกันบ้างก็ตาม)

ทำไมต้อง "ฝัง" theme1.xml ของ 6N5 เข้าไปในไฟล์ใหม่ด้วย (ไม่ใช่แค่ copy สี RGB):
  สีส่วนใหญ่ใน 6N5.xlsx อ้างเป็น theme-color + tint (เช่น theme=3 tint=0.75) ไม่ใช่ RGB ตรงๆ
  ถ้าสร้าง workbook ใหม่ด้วย openpyxl เฉยๆ (ใช้ default Office theme) ตัวเลข theme index เดียวกัน
  จะ resolve ออกมาเป็นสีคนละสีกับที่เห็นใน 6N5.xlsx จริง เพราะ 6N5.xlsx ใช้ custom accent palette
  (accent2=E97132 ส้ม, accent3=196B24 เขียว, accent4=0F9ED5 ฟ้า, ฯลฯ) ไม่ใช่ accent เริ่มต้นของ Excel
  วิธีแก้คือฝัง theme1.xml ชุดเดียวกับ 6N5.xlsx ลง wb.loaded_theme ของ workbook ใหม่ตรงๆ
  (ทดสอบ round-trip แล้วว่า openpyxl เขียน/อ่านกลับมาตรงเดิม) แล้วอ้างสีแบบ theme+tint ต่อได้เป๊ะ

หมายเหตุสำคัญ: 6N5.xlsx เอง "ไม่ได้ commit เข้า git" (อยู่ใน .gitignore บรรทัด *.xlsx เพราะเป็น
ไฟล์ผลลัพธ์/ไฟล์อ้างอิงในเครื่อง) ดังนั้นธีมที่ต้องใช้ตอน deploy จริงต้อง "ฝังไว้ในโค้ด" แบบ base64
ด้านล่างนี้เสมอ — ห้ามพึ่งพาการอ่านไฟล์ 6N5.xlsx จากดิสก์ตอน runtime เพราะอาจไม่มีไฟล์นั้นอยู่เลย
"""
import base64

from openpyxl.cell.text import InlineFont
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.styles.colors import Color

FONT_FAMILY = "Browallia New"

# ── theme1.xml ของ 6N5.xlsx (สกัดจาก xl/theme/theme1.xml จริง, เข้ารหัส base64 ไว้กันไฟล์หาย) ──
# accent1=156082(ฟ้าเข้ม) accent2=E97132(ส้ม) accent3=196B24(เขียว) accent4=0F9ED5(ฟ้าสด)
# accent5=A02B93(ม่วง) accent6=4EA72E(เขียวอ่อน) dk2=0E2841(กรมท่า) lt2=E8E8E8(เทาอ่อน)
_THEME1_XML_B64 = (
    "PD94bWwgdmVyc2lvbj0iMS4wIiBlbmNvZGluZz0iVVRGLTgiIHN0YW5kYWxvbmU9InllcyI/Pg0KPGE6dGhlbWUgeG1sbnM6YT0i"
    "aHR0cDovL3NjaGVtYXMub3BlbnhtbGZvcm1hdHMub3JnL2RyYXdpbmdtbC8yMDA2L21haW4iIG5hbWU9Ik9mZmljZSBUaGVtZSI+"
    "PGE6dGhlbWVFbGVtZW50cz48YTpjbHJTY2hlbWUgbmFtZT0iT2ZmaWNlIj48YTpkazE+PGE6c3lzQ2xyIHZhbD0id2luZG93VGV4"
    "dCIgbGFzdENscj0iMDAwMDAwIi8+PC9hOmRrMT48YTpsdDE+PGE6c3lzQ2xyIHZhbD0id2luZG93IiBsYXN0Q2xyPSJGRkZGRkYi"
    "Lz48L2E6bHQxPjxhOmRrMj48YTpzcmdiQ2xyIHZhbD0iMEUyODQxIi8+PC9hOmRrMj48YTpsdDI+PGE6c3JnYkNsciB2YWw9IkU4"
    "RThFOCIvPjwvYTpsdDI+PGE6YWNjZW50MT48YTpzcmdiQ2xyIHZhbD0iMTU2MDgyIi8+PC9hOmFjY2VudDE+PGE6YWNjZW50Mj48"
    "YTpzcmdiQ2xyIHZhbD0iRTk3MTMyIi8+PC9hOmFjY2VudDI+PGE6YWNjZW50Mz48YTpzcmdiQ2xyIHZhbD0iMTk2QjI0Ii8+PC9h"
    "OmFjY2VudDM+PGE6YWNjZW50ND48YTpzcmdiQ2xyIHZhbD0iMEY5RUQ1Ii8+PC9hOmFjY2VudDQ+PGE6YWNjZW50NT48YTpzcmdi"
    "Q2xyIHZhbD0iQTAyQjkzIi8+PC9hOmFjY2VudDU+PGE6YWNjZW50Nj48YTpzcmdiQ2xyIHZhbD0iNEVBNzJFIi8+PC9hOmFjY2Vu"
    "dDY+PGE6aGxpbms+PGE6c3JnYkNsciB2YWw9IjQ2Nzg4NiIvPjwvYTpobGluaz48YTpmb2xIbGluaz48YTpzcmdiQ2xyIHZhbD0i"
    "OTY2MDdEIi8+PC9hOmZvbEhsaW5rPjwvYTpjbHJTY2hlbWU+PGE6Zm9udFNjaGVtZSBuYW1lPSJPZmZpY2UiPjxhOm1ham9yRm9u"
    "dD48YTpsYXRpbiB0eXBlZmFjZT0iQXB0b3MgRGlzcGxheSIgcGFub3NlPSIwMjExMDAwNDAyMDIwMjAyMDIwNCIvPjxhOmVhIHR5"
    "cGVmYWNlPSIiLz48YTpjcyB0eXBlZmFjZT0iIi8+PGE6Zm9udCBzY3JpcHQ9IkpwYW4iIHR5cGVmYWNlPSLmuLjjgrTjgrfjg4Pj"
    "gq8gTGlnaHQiLz48YTpmb250IHNjcmlwdD0iSGFuZyIgdHlwZWZhY2U9IuunkeydgCDqs6DrlJUiLz48YTpmb250IHNjcmlwdD0i"
    "SGFucyIgdHlwZWZhY2U9Iuetiee6vyBMaWdodCIvPjxhOmZvbnQgc2NyaXB0PSJIYW50IiB0eXBlZmFjZT0i5paw57Sw5piO6auU"
    "Ii8+PGE6Zm9udCBzY3JpcHQ9IkFyYWIiIHR5cGVmYWNlPSJUaW1lcyBOZXcgUm9tYW4iLz48YTpmb250IHNjcmlwdD0iSGViciIg"
    "dHlwZWZhY2U9IlRpbWVzIE5ldyBSb21hbiIvPjxhOmZvbnQgc2NyaXB0PSJUaGFpIiB0eXBlZmFjZT0iVGFob21hIi8+PGE6Zm9u"
    "dCBzY3JpcHQ9IkV0aGkiIHR5cGVmYWNlPSJOeWFsYSIvPjxhOmZvbnQgc2NyaXB0PSJCZW5nIiB0eXBlZmFjZT0iVnJpbmRhIi8+"
    "PGE6Zm9udCBzY3JpcHQ9Ikd1anIiIHR5cGVmYWNlPSJTaHJ1dGkiLz48YTpmb250IHNjcmlwdD0iS2htciIgdHlwZWZhY2U9Ik1v"
    "b2xCb3JhbiIvPjxhOmZvbnQgc2NyaXB0PSJLbmRhIiB0eXBlZmFjZT0iVHVuZ2EiLz48YTpmb250IHNjcmlwdD0iR3VydSIgdHlw"
    "ZWZhY2U9IlJhYXZpIi8+PGE6Zm9udCBzY3JpcHQ9IkNhbnMiIHR5cGVmYWNlPSJFdXBoZW1pYSIvPjxhOmZvbnQgc2NyaXB0PSJD"
    "aGVyIiB0eXBlZmFjZT0iUGxhbnRhZ2VuZXQgQ2hlcm9rZWUiLz48YTpmb250IHNjcmlwdD0iWWlpaSIgdHlwZWZhY2U9Ik1pY3Jv"
    "c29mdCBZaSBCYWl0aSIvPjxhOmZvbnQgc2NyaXB0PSJUaWJ0IiB0eXBlZmFjZT0iTWljcm9zb2Z0IEhpbWFsYXlhIi8+PGE6Zm9u"
    "dCBzY3JpcHQ9IlRoYWEiIHR5cGVmYWNlPSJNViBCb2xpIi8+PGE6Zm9udCBzY3JpcHQ9IkRldmEiIHR5cGVmYWNlPSJNYW5nYWwi"
    "Lz48YTpmb250IHNjcmlwdD0iVGVsdSIgdHlwZWZhY2U9IkdhdXRhbWkiLz48YTpmb250IHNjcmlwdD0iVGFtbCIgdHlwZWZhY2U9"
    "IkxhdGhhIi8+PGE6Zm9udCBzY3JpcHQ9IlN5cmMiIHR5cGVmYWNlPSJFc3RyYW5nZWxvIEVkZXNzYSIvPjxhOmZvbnQgc2NyaXB0"
    "PSJPcnlhIiB0eXBlZmFjZT0iS2FsaW5nYSIvPjxhOmZvbnQgc2NyaXB0PSJNbHltIiB0eXBlZmFjZT0iS2FydGlrYSIvPjxhOmZv"
    "bnQgc2NyaXB0PSJMYW9vIiB0eXBlZmFjZT0iRG9rQ2hhbXBhIi8+PGE6Zm9udCBzY3JpcHQ9IlNpbmgiIHR5cGVmYWNlPSJJc2tv"
    "b2xhIFBvdGEiLz48YTpmb250IHNjcmlwdD0iTW9uZyIgdHlwZWZhY2U9Ik1vbmdvbGlhbiBCYWl0aSIvPjxhOmZvbnQgc2NyaXB0"
    "PSJWaWV0IiB0eXBlZmFjZT0iVGltZXMgTmV3IFJvbWFuIi8+PGE6Zm9udCBzY3JpcHQ9IlVpZ2giIHR5cGVmYWNlPSJNaWNyb3Nv"
    "ZnQgVWlnaHVyIi8+PGE6Zm9udCBzY3JpcHQ9Ikdlb3IiIHR5cGVmYWNlPSJTeWxmYWVuIi8+PGE6Zm9udCBzY3JpcHQ9IkFybW4i"
    "IHR5cGVmYWNlPSJBcmlhbCIvPjxhOmZvbnQgc2NyaXB0PSJCdWdpIiB0eXBlZmFjZT0iTGVlbGF3YWRlZSBVSSIvPjxhOmZvbnQg"
    "c2NyaXB0PSJCb3BvIiB0eXBlZmFjZT0iTWljcm9zb2Z0IEpoZW5nSGVpIi8+PGE6Zm9udCBzY3JpcHQ9IkphdmEiIHR5cGVmYWNl"
    "PSJKYXZhbmVzZSBUZXh0Ii8+PGE6Zm9udCBzY3JpcHQ9Ikxpc3UiIHR5cGVmYWNlPSJTZWdvZSBVSSIvPjxhOmZvbnQgc2NyaXB0"
    "PSJNeW1yIiB0eXBlZmFjZT0iTXlhbm1hciBUZXh0Ii8+PGE6Zm9udCBzY3JpcHQ9Ik5rb28iIHR5cGVmYWNlPSJFYnJpbWEiLz48"
    "YTpmb250IHNjcmlwdD0iT2xjayIgdHlwZWZhY2U9Ik5pcm1hbGEgVUkiLz48YTpmb250IHNjcmlwdD0iT3NtYSIgdHlwZWZhY2U9"
    "IkVicmltYSIvPjxhOmZvbnQgc2NyaXB0PSJQaGFnIiB0eXBlZmFjZT0iUGhhZ3NwYSIvPjxhOmZvbnQgc2NyaXB0PSJTeXJuIiB0"
    "eXBlZmFjZT0iRXN0cmFuZ2VsbyBFZGVzc2EiLz48YTpmb250IHNjcmlwdD0iU3lyaiIgdHlwZWZhY2U9IkVzdHJhbmdlbG8gRWRl"
    "c3NhIi8+PGE6Zm9udCBzY3JpcHQ9IlN5cmUiIHR5cGVmYWNlPSJFc3RyYW5nZWxvIEVkZXNzYSIvPjxhOmZvbnQgc2NyaXB0PSJT"
    "b3JhIiB0eXBlZmFjZT0iTmlybWFsYSBVSSIvPjxhOmZvbnQgc2NyaXB0PSJUYWxlIiB0eXBlZmFjZT0iTWljcm9zb2Z0IFRhaSBM"
    "ZSIvPjxhOmZvbnQgc2NyaXB0PSJUYWx1IiB0eXBlZmFjZT0iTWljcm9zb2Z0IE5ldyBUYWkgTHVlIi8+PGE6Zm9udCBzY3JpcHQ9"
    "IlRmbmciIHR5cGVmYWNlPSJFYnJpbWEiLz48L2E6bWFqb3JGb250PjxhOm1pbm9yRm9udD48YTpsYXRpbiB0eXBlZmFjZT0iQXB0"
    "b3MgTmFycm93IiBwYW5vc2U9IjAyMTEwMDA0MDIwMjAyMDIwMjA0Ii8+PGE6ZWEgdHlwZWZhY2U9IiIvPjxhOmNzIHR5cGVmYWNl"
    "PSIiLz48YTpmb250IHNjcmlwdD0iSnBhbiIgdHlwZWZhY2U9Iua4uOOCtOOCt+ODg+OCryIvPjxhOmZvbnQgc2NyaXB0PSJIYW5n"
    "IiB0eXBlZmFjZT0i66eR7J2AIOqzoOuUlSIvPjxhOmZvbnQgc2NyaXB0PSJIYW5zIiB0eXBlZmFjZT0i562J57q/Ii8+PGE6Zm9u"
    "dCBzY3JpcHQ9IkhhbnQiIHR5cGVmYWNlPSLmlrDntLDmmI7pq5QiLz48YTpmb250IHNjcmlwdD0iQXJhYiIgdHlwZWZhY2U9IkFy"
    "aWFsIi8+PGE6Zm9udCBzY3JpcHQ9IkhlYnIiIHR5cGVmYWNlPSJBcmlhbCIvPjxhOmZvbnQgc2NyaXB0PSJUaGFpIiB0eXBlZmFj"
    "ZT0iVGFob21hIi8+PGE6Zm9udCBzY3JpcHQ9IkV0aGkiIHR5cGVmYWNlPSJOeWFsYSIvPjxhOmZvbnQgc2NyaXB0PSJCZW5nIiB0"
    "eXBlZmFjZT0iVnJpbmRhIi8+PGE6Zm9udCBzY3JpcHQ9Ikd1anIiIHR5cGVmYWNlPSJTaHJ1dGkiLz48YTpmb250IHNjcmlwdD0i"
    "S2htciIgdHlwZWZhY2U9IkRhdW5QZW5oIi8+PGE6Zm9udCBzY3JpcHQ9IktuZGEiIHR5cGVmYWNlPSJUdW5nYSIvPjxhOmZvbnQg"
    "c2NyaXB0PSJHdXJ1IiB0eXBlZmFjZT0iUmFhdmkiLz48YTpmb250IHNjcmlwdD0iQ2FucyIgdHlwZWZhY2U9IkV1cGhlbWlhIi8+"
    "PGE6Zm9udCBzY3JpcHQ9IkNoZXIiIHR5cGVmYWNlPSJQbGFudGFnZW5ldCBDaGVyb2tlZSIvPjxhOmZvbnQgc2NyaXB0PSJZaWlp"
    "IiB0eXBlZmFjZT0iTWljcm9zb2Z0IFlpIEJhaXRpIi8+PGE6Zm9udCBzY3JpcHQ9IlRpYnQiIHR5cGVmYWNlPSJNaWNyb3NvZnQg"
    "SGltYWxheWEiLz48YTpmb250IHNjcmlwdD0iVGhhYSIgdHlwZWZhY2U9Ik1WIEJvbGkiLz48YTpmb250IHNjcmlwdD0iRGV2YSIg"
    "dHlwZWZhY2U9Ik1hbmdhbCIvPjxhOmZvbnQgc2NyaXB0PSJUZWx1IiB0eXBlZmFjZT0iR2F1dGFtaSIvPjxhOmZvbnQgc2NyaXB0"
    "PSJUYW1sIiB0eXBlZmFjZT0iTGF0aGEiLz48YTpmb250IHNjcmlwdD0iU3lyYyIgdHlwZWZhY2U9IkVzdHJhbmdlbG8gRWRlc3Nh"
    "Ii8+PGE6Zm9udCBzY3JpcHQ9Ik9yeWEiIHR5cGVmYWNlPSJLYWxpbmdhIi8+PGE6Zm9udCBzY3JpcHQ9Ik1seW0iIHR5cGVmYWNl"
    "PSJLYXJ0aWthIi8+PGE6Zm9udCBzY3JpcHQ9Ikxhb28iIHR5cGVmYWNlPSJEb2tDaGFtcGEiLz48YTpmb250IHNjcmlwdD0iU2lu"
    "aCIgdHlwZWZhY2U9Iklza29vbGEgUG90YSIvPjxhOmZvbnQgc2NyaXB0PSJNb25nIiB0eXBlZmFjZT0iTW9uZ29saWFuIEJhaXRp"
    "Ii8+PGE6Zm9udCBzY3JpcHQ9IlZpZXQiIHR5cGVmYWNlPSJBcmlhbCIvPjxhOmZvbnQgc2NyaXB0PSJVaWdoIiB0eXBlZmFjZT0i"
    "TWljcm9zb2Z0IFVpZ2h1ciIvPjxhOmZvbnQgc2NyaXB0PSJHZW9yIiB0eXBlZmFjZT0iU3lsZmFlbiIvPjxhOmZvbnQgc2NyaXB0"
    "PSJBcm1uIiB0eXBlZmFjZT0iQXJpYWwiLz48YTpmb250IHNjcmlwdD0iQnVnaSIgdHlwZWZhY2U9IkxlZWxhd2FkZWUgVUkiLz48"
    "YTpmb250IHNjcmlwdD0iQm9wbyIgdHlwZWZhY2U9Ik1pY3Jvc29mdCBKaGVuZ0hlaSIvPjxhOmZvbnQgc2NyaXB0PSJKYXZhIiB0"
    "eXBlZmFjZT0iSmF2YW5lc2UgVGV4dCIvPjxhOmZvbnQgc2NyaXB0PSJMaXN1IiB0eXBlZmFjZT0iU2Vnb2UgVUkiLz48YTpmb250"
    "IHNjcmlwdD0iTXltciIgdHlwZWZhY2U9Ik15YW5tYXIgVGV4dCIvPjxhOmZvbnQgc2NyaXB0PSJOa29vIiB0eXBlZmFjZT0iRWJy"
    "aW1hIi8+PGE6Zm9udCBzY3JpcHQ9Ik9sY2siIHR5cGVmYWNlPSJOaXJtYWxhIFVJIi8+PGE6Zm9udCBzY3JpcHQ9Ik9zbWEiIHR5"
    "cGVmYWNlPSJFYnJpbWEiLz48YTpmb250IHNjcmlwdD0iUGhhZyIgdHlwZWZhY2U9IlBoYWdzcGEiLz48YTpmb250IHNjcmlwdD0i"
    "U3lybiIgdHlwZWZhY2U9IkVzdHJhbmdlbG8gRWRlc3NhIi8+PGE6Zm9udCBzY3JpcHQ9IlN5cmoiIHR5cGVmYWNlPSJFc3RyYW5n"
    "ZWxvIEVkZXNzYSIvPjxhOmZvbnQgc2NyaXB0PSJTeXJlIiB0eXBlZmFjZT0iRXN0cmFuZ2VsbyBFZGVzc2EiLz48YTpmb250IHNj"
    "cmlwdD0iU29yYSIgdHlwZWZhY2U9Ik5pcm1hbGEgVUkiLz48YTpmb250IHNjcmlwdD0iVGFsZSIgdHlwZWZhY2U9Ik1pY3Jvc29m"
    "dCBUYWkgTGUiLz48YTpmb250IHNjcmlwdD0iVGFsdSIgdHlwZWZhY2U9Ik1pY3Jvc29mdCBOZXcgVGFpIEx1ZSIvPjxhOmZvbnQg"
    "c2NyaXB0PSJUZm5nIiB0eXBlZmFjZT0iRWJyaW1hIi8+PC9hOm1pbm9yRm9udD48L2E6Zm9udFNjaGVtZT48YTpmbXRTY2hlbWUg"
    "bmFtZT0iT2ZmaWNlIj48YTpmaWxsU3R5bGVMc3Q+PGE6c29saWRGaWxsPjxhOnNjaGVtZUNsciB2YWw9InBoQ2xyIi8+PC9hOnNv"
    "bGlkRmlsbD48YTpncmFkRmlsbCByb3RXaXRoU2hhcGU9IjEiPjxhOmdzTHN0PjxhOmdzIHBvcz0iMCI+PGE6c2NoZW1lQ2xyIHZh"
    "bD0icGhDbHIiPjxhOmx1bU1vZCB2YWw9IjExMDAwMCIvPjxhOnNhdE1vZCB2YWw9IjEwNTAwMCIvPjxhOnRpbnQgdmFsPSI2NzAw"
    "MCIvPjwvYTpzY2hlbWVDbHI+PC9hOmdzPjxhOmdzIHBvcz0iNTAwMDAiPjxhOnNjaGVtZUNsciB2YWw9InBoQ2xyIj48YTpsdW1N"
    "b2QgdmFsPSIxMDUwMDAiLz48YTpzYXRNb2QgdmFsPSIxMDMwMDAiLz48YTp0aW50IHZhbD0iNzMwMDAiLz48L2E6c2NoZW1lQ2xy"
    "PjwvYTpncz48YTpncyBwb3M9IjEwMDAwMCI+PGE6c2NoZW1lQ2xyIHZhbD0icGhDbHIiPjxhOmx1bU1vZCB2YWw9IjEwNTAwMCIv"
    "PjxhOnNhdE1vZCB2YWw9IjEwOTAwMCIvPjxhOnRpbnQgdmFsPSI4MTAwMCIvPjwvYTpzY2hlbWVDbHI+PC9hOmdzPjwvYTpnc0xz"
    "dD48YTpsaW4gYW5nPSI1NDAwMDAwIiBzY2FsZWQ9IjAiLz48L2E6Z3JhZEZpbGw+PGE6Z3JhZEZpbGwgcm90V2l0aFNoYXBlPSIx"
    "Ij48YTpnc0xzdD48YTpncyBwb3M9IjAiPjxhOnNjaGVtZUNsciB2YWw9InBoQ2xyIj48YTpzYXRNb2QgdmFsPSIxMDMwMDAiLz48"
    "YTpsdW1Nb2QgdmFsPSIxMDIwMDAiLz48YTp0aW50IHZhbD0iOTQwMDAiLz48L2E6c2NoZW1lQ2xyPjwvYTpncz48YTpncyBwb3M9"
    "IjUwMDAwIj48YTpzY2hlbWVDbHIgdmFsPSJwaENsciI+PGE6c2F0TW9kIHZhbD0iMTEwMDAwIi8+PGE6bHVtTW9kIHZhbD0iMTAw"
    "MDAwIi8+PGE6c2hhZGUgdmFsPSIxMDAwMDAiLz48L2E6c2NoZW1lQ2xyPjwvYTpncz48YTpncyBwb3M9IjEwMDAwMCI+PGE6c2No"
    "ZW1lQ2xyIHZhbD0icGhDbHIiPjxhOmx1bU1vZCB2YWw9Ijk5MDAwIi8+PGE6c2F0TW9kIHZhbD0iMTIwMDAwIi8+PGE6c2hhZGUg"
    "dmFsPSI3ODAwMCIvPjwvYTpzY2hlbWVDbHI+PC9hOmdzPjwvYTpnc0xzdD48YTpsaW4gYW5nPSI1NDAwMDAwIiBzY2FsZWQ9IjAi"
    "Lz48L2E6Z3JhZEZpbGw+PC9hOmZpbGxTdHlsZUxzdD48YTpsblN0eWxlTHN0PjxhOmxuIHc9IjEyNzAwIiBjYXA9ImZsYXQiIGNt"
    "cGQ9InNuZyIgYWxnbj0iY3RyIj48YTpzb2xpZEZpbGw+PGE6c2NoZW1lQ2xyIHZhbD0icGhDbHIiLz48L2E6c29saWRGaWxsPjxh"
    "OnByc3REYXNoIHZhbD0ic29saWQiLz48YTptaXRlciBsaW09IjgwMDAwMCIvPjwvYTpsbj48YTpsbiB3PSIxOTA1MCIgY2FwPSJm"
    "bGF0IiBjbXBkPSJzbmciIGFsZ249ImN0ciI+PGE6c29saWRGaWxsPjxhOnNjaGVtZUNsciB2YWw9InBoQ2xyIi8+PC9hOnNvbGlk"
    "RmlsbD48YTpwcnN0RGFzaCB2YWw9InNvbGlkIi8+PGE6bWl0ZXIgbGltPSI4MDAwMDAiLz48L2E6bG4+PGE6bG4gdz0iMjU0MDAi"
    "IGNhcD0iZmxhdCIgY21wZD0ic25nIiBhbGduPSJjdHIiPjxhOnNvbGlkRmlsbD48YTpzY2hlbWVDbHIgdmFsPSJwaENsciIvPjwv"
    "YTpzb2xpZEZpbGw+PGE6cHJzdERhc2ggdmFsPSJzb2xpZCIvPjxhOm1pdGVyIGxpbT0iODAwMDAwIi8+PC9hOmxuPjwvYTpsblN0"
    "eWxlTHN0PjxhOmVmZmVjdFN0eWxlTHN0PjxhOmVmZmVjdFN0eWxlPjxhOmVmZmVjdExzdC8+PC9hOmVmZmVjdFN0eWxlPjxhOmVm"
    "ZmVjdFN0eWxlPjxhOmVmZmVjdExzdC8+PC9hOmVmZmVjdFN0eWxlPjxhOmVmZmVjdFN0eWxlPjxhOmVmZmVjdExzdD48YTpvdXRl"
    "clNoZHcgYmx1clJhZD0iNTcxNTAiIGRpc3Q9IjE5MDUwIiBkaXI9IjU0MDAwMDAiIGFsZ249ImN0ciIgcm90V2l0aFNoYXBlPSIw"
    "Ij48YTpzcmdiQ2xyIHZhbD0iMDAwMDAwIj48YTphbHBoYSB2YWw9IjYzMDAwIi8+PC9hOnNyZ2JDbHI+PC9hOm91dGVyU2hkdz48"
    "L2E6ZWZmZWN0THN0PjwvYTplZmZlY3RTdHlsZT48L2E6ZWZmZWN0U3R5bGVMc3Q+PGE6YmdGaWxsU3R5bGVMc3Q+PGE6c29saWRG"
    "aWxsPjxhOnNjaGVtZUNsciB2YWw9InBoQ2xyIi8+PC9hOnNvbGlkRmlsbD48YTpzb2xpZEZpbGw+PGE6c2NoZW1lQ2xyIHZhbD0i"
    "cGhDbHIiPjxhOnRpbnQgdmFsPSI5NTAwMCIvPjxhOnNhdE1vZCB2YWw9IjE3MDAwMCIvPjwvYTpzY2hlbWVDbHI+PC9hOnNvbGlk"
    "RmlsbD48YTpncmFkRmlsbCByb3RXaXRoU2hhcGU9IjEiPjxhOmdzTHN0PjxhOmdzIHBvcz0iMCI+PGE6c2NoZW1lQ2xyIHZhbD0i"
    "cGhDbHIiPjxhOnRpbnQgdmFsPSI5MzAwMCIvPjxhOnNhdE1vZCB2YWw9IjE1MDAwMCIvPjxhOnNoYWRlIHZhbD0iOTgwMDAiLz48"
    "YTpsdW1Nb2QgdmFsPSIxMDIwMDAiLz48L2E6c2NoZW1lQ2xyPjwvYTpncz48YTpncyBwb3M9IjUwMDAwIj48YTpzY2hlbWVDbHIg"
    "dmFsPSJwaENsciI+PGE6dGludCB2YWw9Ijk4MDAwIi8+PGE6c2F0TW9kIHZhbD0iMTMwMDAwIi8+PGE6c2hhZGUgdmFsPSI5MDAw"
    "MCIvPjxhOmx1bU1vZCB2YWw9IjEwMzAwMCIvPjwvYTpzY2hlbWVDbHI+PC9hOmdzPjxhOmdzIHBvcz0iMTAwMDAwIj48YTpzY2hl"
    "bWVDbHIgdmFsPSJwaENsciI+PGE6c2hhZGUgdmFsPSI2MzAwMCIvPjxhOnNhdE1vZCB2YWw9IjEyMDAwMCIvPjwvYTpzY2hlbWVD"
    "bHI+PC9hOmdzPjwvYTpnc0xzdD48YTpsaW4gYW5nPSI1NDAwMDAwIiBzY2FsZWQ9IjAiLz48L2E6Z3JhZEZpbGw+PC9hOmJnRmls"
    "bFN0eWxlTHN0PjwvYTpmbXRTY2hlbWU+PC9hOnRoZW1lRWxlbWVudHM+PGE6b2JqZWN0RGVmYXVsdHM+PGE6bG5EZWY+PGE6c3BQ"
    "ci8+PGE6Ym9keVByLz48YTpsc3RTdHlsZS8+PGE6c3R5bGU+PGE6bG5SZWYgaWR4PSIyIj48YTpzY2hlbWVDbHIgdmFsPSJhY2Nl"
    "bnQxIi8+PC9hOmxuUmVmPjxhOmZpbGxSZWYgaWR4PSIwIj48YTpzY2hlbWVDbHIgdmFsPSJhY2NlbnQxIi8+PC9hOmZpbGxSZWY+"
    "PGE6ZWZmZWN0UmVmIGlkeD0iMSI+PGE6c2NoZW1lQ2xyIHZhbD0iYWNjZW50MSIvPjwvYTplZmZlY3RSZWY+PGE6Zm9udFJlZiBp"
    "ZHg9Im1pbm9yIj48YTpzY2hlbWVDbHIgdmFsPSJ0eDEiLz48L2E6Zm9udFJlZj48L2E6c3R5bGU+PC9hOmxuRGVmPjwvYTpvYmpl"
    "Y3REZWZhdWx0cz48YTpleHRyYUNsclNjaGVtZUxzdC8+PGE6ZXh0THN0PjxhOmV4dCB1cmk9InswNUE0QzI1Qy0wODVFLTQzNDAt"
    "ODVBMy1BNTUzMUU1MTBEQjJ9Ij48dGhtMTU6dGhlbWVGYW1pbHkgeG1sbnM6dGhtMTU9Imh0dHA6Ly9zY2hlbWFzLm1pY3Jvc29m"
    "dC5jb20vb2ZmaWNlL3RoZW1lbWwvMjAxMi9tYWluIiBuYW1lPSJPZmZpY2UgVGhlbWUiIGlkPSJ7MkUxNDJBMkMtQ0QxNi00MkQ2"
    "LTg3M0EtQzI2RDJBMDUwNkZBfSIgdmlkPSJ7MUJEREZGNTItNkNENi00MEE1LUFCM0MtNjhFQjJGMUU0RDBBfSIvPjwvYTpleHQ+"
    "PC9hOmV4dExzdD48L2E6dGhlbWU+"
)


def get_theme1_xml() -> bytes:
    """theme1.xml ของ 6N5.xlsx เป็น bytes — ใช้ตั้ง wb.loaded_theme ของ workbook ใหม่ก่อน save เสมอ"""
    return base64.b64decode(_THEME1_XML_B64)


def _theme(idx: int, tint: float = 0.0) -> Color:
    return Color(theme=idx, tint=tint, type="theme")


def _rgb(hexcode: str) -> Color:
    return Color(rgb=hexcode if len(hexcode) == 8 else f"FF{hexcode}")


# ── Fill palette ต่อ "role" — ถอดตรงจาก 6N5.xlsx เซลล์จริง (theme index + tint, หรือ rgb ที่เจาะจง) ──
# แต่ละ role คือ "เอกลักษณ์สี" ของหัวข้อ/แถวประเภทนั้นๆ ตามที่เห็นใน 6N5 (Critical IT Load=ฟ้า,
# Transformer=เขียว, Generator=เหลืองทอง, Summary=เหลืองอ่อน ฯลฯ) — ผู้ใช้ยืนยันว่าจุดเด่นของ
# 6N5 คือแต่ละหัวข้อมีสีเป็นเอกลักษณ์ต่างกันชัดเจน ไม่ใช่สีเดียวไล่ทั้งชีท
FILL = {
    "title":            PatternFill("solid", fgColor=_rgb("FFFF00")),   # เหลืองสด — แถบ title ใหญ่สุด (A2)
    "col_header":       PatternFill("solid", fgColor=_theme(3, 0.50)),  # กรมท่าอ่อน — หัวคอลัมน์ Item/Description/Total load
    "normal_band":      PatternFill("solid", fgColor=_rgb("00B050")),   # เขียวสด — แถบ "Normal Operation"
    "fail_band":        PatternFill("solid", fgColor=_theme(5, 0.40)),  # ส้ม — แถบ "X Failure"
    "normal_sub":       PatternFill("solid", fgColor=_theme(6, 0.40)),  # เขียว accent3 — หัวคอลัมน์ย่อยใน Normal
    "fail_sub":         PatternFill("solid", fgColor=_rgb("00B050")),   # เขียวสด — หัวคอลัมน์ย่อยที่ "รอด" ในบล็อก Fail
    "selffail_sub":     PatternFill("solid", fgColor=_rgb("FF0000")),   # แดงสด — หัวคอลัมน์ย่อยของตัวที่ "พัง" เอง
    "data_normal":      PatternFill("solid", fgColor=_theme(5, 0.40)),  # ส้มอ่อน — เซลล์ข้อมูลใต้ Normal
    "data_fail":        PatternFill("solid", fgColor=_theme(9, 0.80)),  # เขียวอ่อน accent6 — เซลล์ข้อมูลใต้ Fail (รอด)
    "data_selffail":    PatternFill("solid", fgColor=_theme(0, -0.15)), # เทา — เซลล์คอลัมน์ของตัวที่พังเอง (skip)
    "spacer":           PatternFill("solid", fgColor=_theme(8, 0.80)),  # ม่วงอ่อน accent5 — คอลัมน์คั่นระหว่าง scenario
    "section_header":   PatternFill("solid", fgColor=_theme(7, 0.60)),  # ฟ้าสด accent4 — หัวข้อหลัก (Critical IT Load ฯลฯ)
    "phase_header":     PatternFill("solid", fgColor=_theme(7, 0.80)),  # ฟ้าอ่อน accent4 — หัวข้อย่อยระดับ phase/group
    "subtotal":         PatternFill("solid", fgColor=_theme(3, 0.90)),  # กรมท่าอ่อนมาก — แถว subtotal ทั่วไป
    "capacity":         PatternFill("solid", fgColor=_theme(3, 0.75)),  # กรมท่าอ่อน — แถว Capacity/Max Connected Load
    "total_connected":  PatternFill("solid", fgColor=_theme(3, 0.60)),  # กรมท่ากลาง — แถว Total Connected Load (เน้นสุด)
    "summary_highlight":PatternFill("solid", fgColor=_rgb("FFFFCC")),   # เหลืองอ่อน — Summary of UPS/HVAC Losses
    "transformer":      PatternFill("solid", fgColor=_theme(6, 0.40)),  # เขียว accent3 — Capacity of Transformer
    "transformer_util": PatternFill("solid", fgColor=_theme(6, 0.60)),  # เขียวอ่อนกว่า — Utilization ของ Transformer
    "generator":        PatternFill("solid", fgColor=_rgb("FFC000")),   # เหลืองทอง — Capacity of Generator
    "generator_util":   PatternFill("solid", fgColor=_rgb("FFFF00")),   # เหลืองสด — Utilization ของ Generator
    "band_alt":         PatternFill("solid", fgColor=_theme(3, 0.95)),  # ลายทางสลับอ่อนมาก สำหรับ data row คู่/คี่
}

# ── Font palette (sheet กลุ่ม) — ผู้ใช้ยืนยัน 2026-09: ให้ขนาดตัวอักษร "24 ทุกอย่างเลย"
# (เหมือนกด Ctrl+A เลือกทั้งชีทแล้วตั้ง font size เดียวเท่ากันหมด — ไม่ไล่ระดับใหญ่เล็กตาม hierarchy
# แบบที่ 6N5 ทำจริง (6N5 title 36 / body 24 / footnote 26) เจาะจงว่าต้องการ 24 เท่ากันทุก role) ──
GROUP_FONT_SIZE = 24
FONT = {
    "title":       Font(name=FONT_FAMILY, size=GROUP_FONT_SIZE, bold=True),
    "header":      Font(name=FONT_FAMILY, size=GROUP_FONT_SIZE, bold=True),
    "header_wh":   Font(name=FONT_FAMILY, size=GROUP_FONT_SIZE, bold=True, color=_rgb("FFFFFF")),
    "body":        Font(name=FONT_FAMILY, size=GROUP_FONT_SIZE),
    "body_bold":   Font(name=FONT_FAMILY, size=GROUP_FONT_SIZE, bold=True),
    "fail":        Font(name=FONT_FAMILY, size=GROUP_FONT_SIZE, bold=True, color=_rgb("C00000")),
    "fail_receive":Font(name=FONT_FAMILY, size=GROUP_FONT_SIZE, bold=True, color=_rgb("375623")),
    "footnote":    Font(name=FONT_FAMILY, size=GROUP_FONT_SIZE, bold=True, italic=True),
}

# ── InlineFont สำหรับ rich text ในเซลล์ FAIL (ตัวอักษรพัง=แดง, ตัวอักษรที่รับโหลดแทน=เขียว) ──
# ต้องใช้ InlineFont (ไม่ใช่ Font ธรรมดา) เพราะเป็นการผสมหลายสีในเซลล์เดียวกัน (rich text run)
RICH_FAIL = InlineFont(rFont=FONT_FAMILY, sz=GROUP_FONT_SIZE, b=True, color="FFC00000")
RICH_RECEIVE = InlineFont(rFont=FONT_FAMILY, sz=GROUP_FONT_SIZE, b=True, color="FF375623")
RICH_ARROW = InlineFont(rFont=FONT_FAMILY, sz=GROUP_FONT_SIZE, b=True, color="FF000000")

# ── Border ต้นแบบ (สีอัตโนมัติ/ดำ เหมือน 6N5 — ไม่ใช้สี theme กับเส้นขอบ) ──────
_SIDE_MEDIUM = Side(style="medium", color="000000")
_SIDE_THIN = Side(style="thin", color="000000")
_SIDE_HAIR = Side(style="hair", color="000000")

BORDER_OUTER = Border(left=_SIDE_MEDIUM, right=_SIDE_MEDIUM, top=_SIDE_MEDIUM, bottom=_SIDE_MEDIUM)
BORDER_DATA_ROW = Border(left=_SIDE_MEDIUM, right=_SIDE_MEDIUM, top=_SIDE_HAIR, bottom=_SIDE_HAIR)
BORDER_SUBTOTAL_ROW = Border(left=_SIDE_MEDIUM, right=_SIDE_MEDIUM, top=_SIDE_THIN, bottom=_SIDE_THIN)
BORDER_INNER_THIN = Border(left=_SIDE_THIN, right=_SIDE_THIN, top=_SIDE_THIN, bottom=_SIDE_THIN)
BORDER_INNER_HAIR = Border(left=_SIDE_HAIR, right=_SIDE_HAIR, top=_SIDE_HAIR, bottom=_SIDE_HAIR)

CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
CENTER_NOWRAP = Alignment(horizontal="center", vertical="center", wrap_text=False)  # เหมือน 6N5: ข้อความ "A → B" ไม่ wrap ปล่อยล้นได้
LEFT = Alignment(horizontal="left", vertical="center")
RIGHT = Alignment(horizontal="right", vertical="center")

# ── ความกว้างคอลัมน์ (ถอดจาก 6N5.xlsx) ─────────────────────────────────────
COL_WIDTH_ITEM = 8.71       # A: เลขข้อ
COL_WIDTH_DESC = 82.0       # B: คำอธิบาย — กว้างพอไม่ตัดคำ label ยาวสุด
COL_WIDTH_KW = 12.43        # C: Load (kW)
COL_WIDTH_PF = 9.29         # D: PF
COL_WIDTH_KVA = 13.86       # E: kVA
COL_WIDTH_DATA_FIRST = 13.86  # คอลัมน์แรกของแต่ละ scenario block
COL_WIDTH_DATA = 12.71        # คอลัมน์ข้อมูลอื่นๆ ใน scenario block
COL_WIDTH_SPACER = 1.43       # คอลัมน์คั่นระหว่าง scenario

ROW_HEIGHT = 30.0
ROW_HEIGHT_TITLE = 34.0
ROW_HEIGHT_HEADER = 31.0

NUMFMT_KW = "#,##0.00"
NUMFMT_KVA = "#,##0.00"
NUMFMT_SUBTOTAL = "#,##0.0"
NUMFMT_MAX = "#,##0"
NUMFMT_PCT = "0.0%"


def apply_theme(wb) -> None:
    """ฝัง theme1.xml ของ 6N5.xlsx ลง workbook ใหม่ — ต้องเรียกก่อน wb.save() เสมอ
    ไม่งั้นสี theme+tint ใน FILL ด้านบนจะ resolve ผิดเป็นสี default ของ Excel"""
    wb.loaded_theme = get_theme1_xml()


def build_fail_arrow_text(faulted_label: str, survivor_labels: list):
    """
    สร้างข้อความ rich text แบบ "A → B" (หรือ "A → B,C,D" ถ้าแถวนั้นเป็น 4-source) สำหรับเซลล์
    self-fail ของแถวข้อมูลที่ UPS ตัวนี้ "อยู่ในสาย pair" ของแถวนั้นจริงๆ (ยืนยันโดยผู้ใช้ 2026-09:
    อยากเห็นว่าโหลดที่พังไปวิ่งไปที่ตัวไหนต่อ แบบเดียวกับที่ 6N5.xlsx ใช้ "A → B") — ตัวที่พัง (faulted)
    เป็นตัวอักษรสีแดง ตัวที่รับโหลดแทน (survivor) เป็นตัวอักษรสีเขียว ส่วนลูกศรตรงกลางเป็นสีดำเฉยๆ
    (ต่างจาก 6N5 ซึ่งทั้งข้อความเป็นสีดำล้วน — ผู้ใช้ขอเพิ่มสีแยกสองฝั่งเองโดยเฉพาะ)

    survivor_labels ว่างเปล่าไม่ควรเกิดขึ้น (compute_fault_loads รับประกันว่าถ้า faulted อยู่ใน
    row['pair'] จะต้องมีสมาชิกที่เหลืออย่างน้อย 1 ตัวเสมอ เพราะ pair มีอย่างน้อย 2 ตัวอักษร)
    """
    from openpyxl.cell.rich_text import CellRichText, TextBlock
    receivers = ",".join(survivor_labels)
    return CellRichText(
        TextBlock(RICH_FAIL, faulted_label),
        TextBlock(RICH_ARROW, " → "),
        TextBlock(RICH_RECEIVE, receivers),
    )
